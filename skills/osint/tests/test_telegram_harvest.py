"""telegram-harvest: rule extraction, geo clues, confidence, workbook + CSV offline mode."""
import csv
import json
import sys
from pathlib import Path

import pytest

SKILL = Path(__file__).resolve().parents[2] / "telegram-harvest"
sys.path.insert(0, str(SKILL / "scripts"))
import extract  # noqa: E402
import geo_clues  # noqa: E402
import harvest  # noqa: E402

KW = extract.load_keywords()
CO = geo_clues.load_countries()

POST = ("Selling verified AdSense accounts, US acc, PIN verified, payment attached. "
        "In stock: 12 pcs. Price $150 each. Contact @seller_one or t.me/seller_chan. Pay USDT.")


def test_extract_listing_fields():
    f = extract.extract_listing(POST, "", KW)
    assert f["account_type"] == "AdSense"
    assert f["verified"] and f["pin"] and f["payment_attached"]
    assert f["cost"] == "$150"
    assert f["quantity"] == "12"
    assert f["payment_method"] == "usdt"
    assert "@seller_one" in f["handles"]
    assert any("t.me/seller_chan" in link for link in f["links"])


def test_non_listing_and_russian_terms():
    assert not extract.is_listing("hello world, nothing for sale", KW)
    ru = "Продам аккаунты admob, верифицированные, цена 300 руб"
    f = extract.extract_listing(ru, "", KW)
    assert f["account_type"] == "AdMob" and f["verified"] and f["cost"]


def test_geo_clues_never_decide():
    clues = geo_clues.gather_clues(POST, "", "English", CO)
    assert "United States" in clues["country_words"]
    assert set(clues) == {"language", "country_words"}  # no verdict field
    assert geo_clues.phone_country("+44 7700 900000") == "United Kingdom"


def test_build_record_and_confidence():
    rec = harvest.build_record({"message_id": "77", "date": "2026-10-01", "text": POST},
                               "@seller_chan", KW, CO, "2026-10-10", translate=False)
    assert rec["url"] == "https://t.me/seller_chan/77"
    assert rec["confidence"] == "high" and rec["needs_review"] is False
    assert rec["tier"] == "T3"
    assert harvest.build_record({"text": "just chatting"}, "@x", KW, CO, "d", translate=False) is None
    assert harvest.message_url("t.me/+AbCdEf", "5") == "t.me/+AbCdEf"  # invite links have no message url


def test_from_csv_end_to_end(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    src = tmp_path / "messages.csv"
    with src.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["message_id", "date", "sender", "text"])
        w.writeheader()
        w.writerow({"message_id": "1", "date": "2026-10-01", "sender": "s", "text": POST})
        w.writerow({"message_id": "2", "date": "2026-10-01", "sender": "s", "text": "good morning"})
        w.writerow({"message_id": "3", "date": "2026-10-02", "sender": "s", "text": "YPP channel for sale"})
    out = tmp_path / "run" / "sheet.xlsx"
    rc = harvest.main(["--from-csv", str(src), "--channel", "@seller_chan", "--no-translate",
                       "--out", str(out)])
    assert rc == 0
    ws = load_workbook(out)["Account Procurement"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["Date of Finding", "URL"] and "» Clue: TG Handles" in header
    assert ws.max_row == 3  # header + 2 listings
    geo_col = header.index("Geo of Acc.")
    assert all(row[geo_col].value in (None, "") for row in ws.iter_rows(min_row=2))
    listings = [json.loads(line) for line in (tmp_path / "run" / "listings.jsonl").read_text().splitlines()]
    assert [entry["message_id"] for entry in listings] == ["1", "3"]
    tsv = (tmp_path / "run" / "urls.tsv").read_text().splitlines()
    assert tsv[1].split("\t")[:3] == ["https://t.me/seller_chan/1", "tg-seller_chan-1", "telegram"]


def test_max_messages_bound(tmp_path):
    pytest.importorskip("openpyxl")
    src = tmp_path / "m.csv"
    with src.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["message_id", "date", "sender", "text"])
        w.writeheader()
        for i in range(5):
            w.writerow({"message_id": str(i), "date": "", "sender": "", "text": POST})
    out = tmp_path / "r" / "s.xlsx"
    harvest.main(["--from-csv", str(src), "--channel", "@c", "--no-translate", "--out", str(out),
                  "--max-messages", "2"])
    assert len((tmp_path / "r" / "listings.jsonl").read_text().splitlines()) == 2
