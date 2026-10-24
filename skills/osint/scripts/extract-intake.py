#!/usr/bin/env python3
"""extract-intake.py -- canonical client-document extraction for Q investigations.

Pairs with ingest-client-document.sh. That script registers each original file as
an EV-NNNN item with SHA-256 + chain of custody. THIS script does the verbatim
content extraction + full OCR that the evidence protocol requires, deterministically
(no network, no LLM -- same input always yields the same output).

For every file in a case's investigation/intake/, writes:
  investigation/evidence/extracted/<stem>/
    text.md        -- verbatim text (PDF text layer / DOCX paragraphs+tables / DOC via textutil)
    images/        -- embedded images (DOCX/DOC) and full-page renders (PDF)
    ocr.md         -- tesseract OCR (multi-PSM, best-by-alnum) of every image/page
    manifest.json  -- per-file extraction stats for verification

    iocs.json      -- deterministic IOCs / ATT&CK ids / gazetteer hits (lib/ioc_extract.py,
                      theia port) found in text.md + ocr.md, each with exact offsets
    entities.json  -- emails / handles / phones / wallets / people + org candidates
                      (lib/entity_extract.py, kipi port), each with exact offsets
    exif.json      -- EXIF of every image under images/ + attachments/ (exiftool, else
                      Pillow, else recorded as unavailable). Not written when no images.
    ner.json       -- only with --ner: optional model NER (GLiNER / langextract / spaCy)

Side-file offsets are Python str indices into the named source file (text.md or
ocr.md): source[start:end] == the recorded span. Side-file failures are reported in
the manifest and on stderr; they never fail the verbatim extraction itself.

Handlers: .pdf .docx .doc .csv .eml .md .json (Telegram export / generic JSON) .xlsx
          .png .jpg .jpeg .tif .tiff
Requires on PATH: pdftotext, pdftoppm, tesseract. Python: python-docx + Pillow
(only needed when DOCX/DOC/image inputs are present; imported lazily); openpyxl for
.xlsx (lazy; a clear FAIL row if missing). Optional: exiftool on PATH.

Usage:
  python3 extract-intake.py --case <case-folder> [--intake <dir>] [--ner BACKEND]
    --case    investigations/<case-folder>; defaults to <workspace>/.active-case
    --intake  override the intake dir (default: <case>/investigation/intake)
    --ner     gliner | langextract | spacy  (optional model NER -> ner.json).
              gliner: local model (pip install gliner; weights download once from
              Hugging Face, model via $GLINER_MODEL). langextract: SENDS TEXT TO
              GOOGLE GEMINI, needs $LANGEXTRACT_API_KEY or $GEMINI_API_KEY, skipped
              otherwise. spacy: EntityRuler regex patterns (IBAN, wallets,
              passport/ID-like) + en_core_web_sm PERSON/ORG when installed.

Exit codes: 0 all files extracted; 1 setup error (no case/intake/files);
            2 one or more files failed extraction (others still processed).
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

EXTRACTOR_NAME = "extract-intake.py"

# workspace = workspace (script lives in skills/osint/scripts/)
WORKSPACE = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(Path(__file__).resolve().parent / "lib"))
from active_case import resolve_case_dir  # noqa: E402
import entity_extract  # noqa: E402
import ioc_extract  # noqa: E402

PSM_MODES = ["3", "6", "11"]
PSM_LABELS = {
    "3": "PSM 3 (auto)",
    "6": "PSM 6 (uniform block)",
    "11": "PSM 11 (sparse text)",
    "up4x": "upscale 4x + binarize + PSM 11",
}
RENDER_DPI = "200"
ALNUM_RE = re.compile(r"[A-Za-z0-9]")
TRAILING_NUM_RE = re.compile(r"(\d+)(?=\D*$)")


# ---------- small shared helpers ----------

def run_cmd(cmd: list) -> str:
    result = subprocess.run(cmd, capture_output=True, text=True, check=False)
    return result.stdout


def alnum_count(text: str) -> int:
    return len(ALNUM_RE.findall(text))


def write_text_md(out_dir: Path, src: Path, text: str, note: str) -> None:
    body = [f"# {src.name}", "", f"_Extraction: {note}_", "",
            "## Verbatim text", "", text.rstrip(), ""]
    (out_dir / "text.md").write_text("\n".join(body), encoding="utf-8")


def sort_pages(paths: list) -> list:
    def key(p: Path) -> int:
        match = TRAILING_NUM_RE.search(p.stem)
        return int(match.group(1)) if match else 0
    return sorted(paths, key=key)


# ---------- OCR (multi-PSM, deterministic) ----------

def tesseract_psm(image: Path, psm: str) -> str:
    result = subprocess.run(
        ["tesseract", str(image), "-", "-l", "eng", "--psm", psm],
        capture_output=True, text=True, check=False,
    )
    return result.stdout.rstrip() if result.returncode == 0 else ""


def upscale_and_binarize(image: Path):
    from PIL import Image
    try:
        img = Image.open(image)
    except Exception:
        return None
    if img.width < 100 or img.height < 100:
        return None
    upscaled = img.resize((img.width * 4, img.height * 4), Image.LANCZOS).convert("L")
    binarized = upscaled.point(lambda p: 255 if p > 140 else 0)
    out_path = image.with_suffix(".up4x.png")
    binarized.save(out_path)
    return out_path


def ocr_image(image: Path):
    """Return (best_text, best_label). Tries 3 PSM modes, upscales if all empty."""
    outputs = {psm: tesseract_psm(image, psm) for psm in PSM_MODES}
    if all(alnum_count(outputs[p]) == 0 for p in PSM_MODES):
        upscaled = upscale_and_binarize(image)
        outputs["up4x"] = tesseract_psm(upscaled, "11") if upscaled else ""
    else:
        outputs["up4x"] = ""
    order = PSM_MODES + ["up4x"]
    best = max(order, key=lambda p: (alnum_count(outputs[p]), -order.index(p)))
    return outputs[best], best


def ocr_images(out_dir: Path, src: Path, images: list, header: str) -> int:
    if not images:
        return 0
    lines = [f"# {src.name} -- OCR ({header})", ""]
    total_chars = 0
    for img in images:
        best_text, best_label = ocr_image(img)
        total_chars += len(best_text)
        lines += [f"## {img.name}  ({PSM_LABELS[best_label]})", "",
                  "```", best_text or "(no text detected)", "```", ""]
    (out_dir / "ocr.md").write_text("\n".join(lines), encoding="utf-8")
    return total_chars


# ---------- DOCX helpers ----------

def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    lines = ["## Paragraphs", ""]
    for para in doc.paragraphs:
        if para.text.strip():
            lines += [para.text, ""]
    table_no = 0
    for table in doc.tables:
        table_no += 1
        lines += [f"## Table {table_no}", ""]
        for row in table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("")
    return "\n".join(lines)


def docx_images(path: Path, images_dir: Path) -> list:
    from docx import Document
    images_dir.mkdir(parents=True, exist_ok=True)
    doc = Document(str(path))
    saved = []
    for rel in doc.part.related_parts.values():
        content_type = getattr(rel, "content_type", "") or ""
        if not content_type.startswith("image/"):
            continue
        ext = content_type.split("/")[-1].split("+")[0]
        if ext == "jpeg":
            ext = "jpg"
        img_path = images_dir / f"image_{len(saved) + 1:03d}.{ext}"
        img_path.write_bytes(rel.blob)
        saved.append(img_path)
    return saved


# ---------- per-type handlers ----------

def extract_pdf(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    images_dir = out / "images"
    images_dir.mkdir(exist_ok=True)

    text = run_cmd(["pdftotext", "-layout", str(path), "-"])
    write_text_md(out, path, text, "PDF text layer via pdftotext -layout")

    subprocess.run(["pdftoppm", "-png", "-r", RENDER_DPI, str(path), str(images_dir / "page")],
                   check=False)
    pages = sort_pages(list(images_dir.glob("page*.png")))
    ocr_images(out, path, pages, f"full-page renders, {len(pages)} pages")
    return {"text_chars": len(text), "images": len(pages)}


def extract_docx(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    text = docx_text(path)
    write_text_md(out, path, text, "DOCX paragraphs + tables via python-docx")
    images = docx_images(path, out / "images")
    ocr_images(out, path, images, "DOCX embedded images")
    return {"text_chars": len(text), "images": len(images)}


def extract_doc(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    text = run_cmd(["textutil", "-convert", "txt", "-stdout", str(path)])
    write_text_md(out, path, text, "legacy .doc text via textutil")

    converted = out / f"{path.stem}.converted.docx"
    subprocess.run(["textutil", "-convert", "docx", "-output", str(converted), str(path)],
                   check=False)
    images = docx_images(converted, out / "images") if converted.exists() else []
    ocr_images(out, path, images, "DOC embedded images (textutil -> docx)")
    return {"text_chars": len(text), "images": len(images)}


def extract_image(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    images_dir = out / "images"
    images_dir.mkdir(exist_ok=True)
    dest = images_dir / path.name
    shutil.copy2(path, dest)
    write_text_md(out, path, "(binary image -- OCR only, see ocr.md)", "image file")
    ocr_images(out, path, [dest], "image OCR")
    return {"text_chars": 0, "images": 1}


def extract_csv(path: Path, out: Path) -> dict:
    import csv
    out.mkdir(parents=True, exist_ok=True)
    lines = [f"# {path.name}", "", "_Extraction: CSV rows as structured text (csv module)_", ""]
    row_count = 0
    headers = []
    try:
        with path.open(encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.DictReader(fh)
            headers = reader.fieldnames or []
            for i, row in enumerate(reader):
                non_empty = {k: v for k, v in row.items() if k and v and v.strip()}
                if not non_empty:
                    continue
                lines.append(f"## Row {i + 1}")
                lines.append("")
                for k, v in non_empty.items():
                    lines.append(f"**{k}:** {v.strip()}")
                lines.append("")
                row_count += 1
    except Exception as exc:
        lines.append(f"_(CSV parse error: {exc})_")
    text = "\n".join(lines)
    (out / "text.md").write_text(text, encoding="utf-8")
    # No images in CSV; write empty ocr.md for consistency
    (out / "ocr.md").write_text(f"# {path.name} -- OCR\n\n_(CSV file -- no images to OCR)_\n",
                                encoding="utf-8")
    return {"text_chars": len(text), "images": 0, "rows": row_count, "columns": len(headers)}


def extract_eml(path: Path, out: Path) -> dict:
    """Extract an .eml email: headers, body (plain text, HTML fallback),
    inline images (OCR'd via the shared multi-PSM pipeline), attachments
    saved verbatim, and any .docx attachment additionally extracted into
    its own subdirectory. Ported 2026-07-21 from case-025/case-028's
    independently-written extract_intake.py -- this capability (email
    extraction) did not exist in this module until now."""
    import email
    import email.policy

    out.mkdir(parents=True, exist_ok=True)
    images_dir = out / "images"
    attachments_dir = out / "attachments"
    images_dir.mkdir(exist_ok=True)
    attachments_dir.mkdir(exist_ok=True)

    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)

    lines = [f"# {path.name}", "", "## Headers", "",
             f"- **From:** {msg.get('From', '(none)')}",
             f"- **To:** {msg.get('To', '(none)')}",
             f"- **CC:** {msg.get('Cc', '(none)')}",
             f"- **Date:** {msg.get('Date', '(none)')}",
             f"- **Subject:** {msg.get('Subject', '(none)')}",
             "", "## Body", ""]

    body_parts, image_paths, attachment_paths = [], [], []
    for part in msg.walk():
        content_type = part.get_content_type()
        disposition = str(part.get("Content-Disposition", ""))

        if content_type == "text/plain" and "attachment" not in disposition:
            try:
                body_parts.append(part.get_content())
            except Exception as exc:
                body_parts.append(f"(decode error: {exc})")
        elif content_type == "text/html" and "attachment" not in disposition and not body_parts:
            try:
                html = part.get_content()
                body_parts.append(f"[HTML body -- first 4000 chars]\n{html[:4000]}")
            except Exception as exc:
                body_parts.append(f"(HTML decode error: {exc})")
        elif content_type.startswith("image/"):
            ext = content_type.split("/")[-1].split("+")[0]
            ext = "jpg" if ext == "jpeg" else ext
            img_path = images_dir / f"image_{len(image_paths) + 1:03d}.{ext}"
            try:
                img_path.write_bytes(part.get_payload(decode=True))
                image_paths.append(img_path)
            except Exception as exc:
                print(f"  [warn] inline image failed: {exc}", file=sys.stderr)
        elif "attachment" in disposition or content_type not in (
            "text/plain", "text/html", "multipart/mixed",
            "multipart/alternative", "multipart/related", "multipart/signed",
        ):
            filename = part.get_filename()
            if filename:
                att_path = attachments_dir / filename
                try:
                    att_path.write_bytes(part.get_payload(decode=True))
                    attachment_paths.append(att_path)
                except Exception as exc:
                    print(f"  [warn] attachment {filename} failed: {exc}", file=sys.stderr)

    lines.extend(body_parts if body_parts else ["(no plain-text body found)"])
    lines.append("")
    if attachment_paths:
        lines += ["## Attachments", ""]
        lines += [f"- `{a.name}` -> `attachments/{a.name}`" for a in attachment_paths]
        lines.append("")
    if image_paths:
        lines += ["## Inline Images", ""]
        lines += [f"- `{i.name}` -> `images/{i.name}`" for i in image_paths]
        lines.append("")
    (out / "text.md").write_text("\n".join(lines), encoding="utf-8")

    ocr_images(out, path, image_paths, "EML inline images")

    for att in attachment_paths:
        if att.suffix.lower() == ".docx":
            try:
                extract_docx(att, attachments_dir / f"{att.stem}_extracted")
            except Exception as exc:
                print(f"  [warn] nested docx extract failed for {att.name}: {exc}", file=sys.stderr)

    return {"text_chars": sum(len(b) for b in body_parts), "images": len(image_paths),
            "attachments": len(attachment_paths)}


def extract_md(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    text = path.read_text(encoding="utf-8", errors="replace")
    write_text_md(out, path, text, "Markdown file, copied verbatim")
    (out / "ocr.md").write_text(f"# {path.name} -- OCR\n\n_(Markdown file -- no images to OCR)_\n",
                                encoding="utf-8")
    return {"text_chars": len(text), "images": 0}


def _write_no_ocr(out: Path, src: Path, kind: str) -> None:
    (out / "ocr.md").write_text(f"# {src.name} -- OCR\n\n_({kind} -- no images to OCR)_\n",
                                encoding="utf-8")


# ---------- Telegram / generic JSON (port of kipi ingest/telegram.py) ----------

TELEGRAM_META_KEYS = ("date", "from", "from_id", "author", "sender", "username")


def _telegram_msg(item) -> str:
    if not isinstance(item, dict):
        return str(item)
    parts = [f"{key}={item[key]}" for key in TELEGRAM_META_KEYS if item.get(key)]
    text = item.get("text") or item.get("message") or item.get("content") or ""
    if isinstance(text, list):
        # Telegram splits one message into contiguous text/entity parts; join with ""
        # so the message stays verbatim (kipi joined with " ", inserting spaces).
        text = "".join((t.get("text", "") if isinstance(t, dict) else str(t)) for t in text)
    if parts or text:
        prefix = f"[{' '.join(parts)}] " if parts else ""
        return f"{prefix}{text}".strip()
    return ""


def telegram_text(data) -> tuple:
    """Return (text, message_count). Native channel export ({name, messages: [...]}),
    kipi harvest list, or a single message dict."""
    msgs = []
    count = 0
    if isinstance(data, list):
        msgs = [_telegram_msg(item) for item in data]
        count = len(data)
    elif isinstance(data, dict) and isinstance(data.get("messages"), list):
        channel = data.get("name") or data.get("title") or "unknown_channel"
        msgs = [f"# Channel: {channel}\n"] + [_telegram_msg(i) for i in data["messages"]]
        count = len(data["messages"])
    elif isinstance(data, dict):
        msgs = [_telegram_msg(data)]
        count = 1
    return "\n".join(m for m in msgs if m), count


def extract_json(path: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        write_text_md(out, path, raw, f"JSON parse failed ({exc}); raw text copied verbatim")
        _write_no_ocr(out, path, "JSON file")
        return {"text_chars": len(raw), "images": 0, "messages": 0}
    text, count = telegram_text(data)
    write_text_md(out, path, text,
                  "JSON as Telegram-style messages (kipi telegram.py format); "
                  "original JSON stays in intake/")
    _write_no_ocr(out, path, "JSON file")
    return {"text_chars": len(text), "images": 0, "messages": count}


# ---------- XLSX (openpyxl, lazy) ----------

def extract_xlsx(path: Path, out: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed: pip install openpyxl") from exc
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines = [f"# {path.name}", "", "_Extraction: XLSX rows as structured text (openpyxl)_", ""]
    row_count = 0
    sheets = 0
    for ws in wb.worksheets:
        sheets += 1
        lines += [f"## Sheet: {ws.title}", ""]
        headers = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = ["" if v is None else str(v).strip() for v in row]
            if i == 0:
                headers = [h or f"col{j + 1}" for j, h in enumerate(cells)]
                continue
            pairs = [(headers[j] if j < len(headers) else f"col{j + 1}", v)
                     for j, v in enumerate(cells) if v]
            if not pairs:
                continue
            lines += [f"### Row {i}", ""] + [f"**{k}:** {v}" for k, v in pairs] + [""]
            row_count += 1
    wb.close()
    text = "\n".join(lines)
    (out / "text.md").write_text(text, encoding="utf-8")
    _write_no_ocr(out, path, "XLSX file")
    return {"text_chars": len(text), "images": 0, "rows": row_count, "sheets": sheets}


# ---------- side files: iocs.json / entities.json / exif.json / ner.json ----------

SIDE_SOURCES = ("text.md", "ocr.md")
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".gif", ".webp", ".heic", ".bmp"}
SIDE_GENERATOR = "extract-intake.py side files v1"


def read_side_sources(out: Path) -> dict:
    return {name: (out / name).read_text(encoding="utf-8")
            for name in SIDE_SOURCES if (out / name).is_file()}


def _tag(items: list, source: str) -> list:
    return [{**it, "source_file": source} for it in items]


def write_iocs_and_entities(out: Path, sources: dict) -> dict:
    iocs, attack, gaz, ents = [], [], [], []
    for name, text in sources.items():
        res = ioc_extract.extract_all(text)
        iocs += _tag(res["iocs"], name)
        attack += _tag(res["attack_ids"], name)
        for group in res["entities"].values():
            gaz += _tag(group, name)
        ents += _tag(entity_extract.extract_dicts(text), name)
    header = {"generator": SIDE_GENERATOR, "offset_unit": "python str index (code point)",
              "sources": sorted(sources), "evidence_tier": "T3 (automated extraction)"}
    (out / "iocs.json").write_text(json.dumps(
        {**header, "iocs": iocs, "attack_ids": attack, "gazetteer_entities": gaz},
        indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "entities.json").write_text(json.dumps(
        {**header, "entities": ents}, indent=2, ensure_ascii=False), encoding="utf-8")
    return {"iocs": len(iocs), "attack_ids": len(attack),
            "gazetteer_entities": len(gaz), "entities": len(ents)}


def find_images(out: Path, ext: str) -> list:
    found = []
    for sub in ("images", "attachments"):
        base = out / sub
        if not base.is_dir():
            continue
        for p in sorted(base.rglob("*")):
            if not p.is_file() or p.suffix.lower() not in IMAGE_SUFFIXES:
                continue
            if p.name.endswith(".up4x.png"):
                continue  # our own OCR upscale artifact
            if ext == ".pdf" and p.parent == out / "images" and p.name.startswith("page"):
                continue  # our own pdftoppm render, carries no source metadata
            found.append(p)
    return found


def exif_exiftool(images: list, out: Path) -> dict:
    result = subprocess.run(["exiftool", "-j", "-n", "-G1", *[str(p) for p in images]],
                            capture_output=True, text=True, check=False)
    if not result.stdout.strip():
        raise RuntimeError(f"exiftool produced no output: {result.stderr.strip()[:300]}")
    records = json.loads(result.stdout)
    by_name = {}
    for rec in records:
        src = Path(rec.get("SourceFile", ""))
        key = str(src.relative_to(out)) if src.is_absolute() and out in src.parents else str(src)
        by_name[key] = rec
    return by_name


def _jsonable(value):
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    if isinstance(value, (int, float, str)) or value is None:
        return value
    if isinstance(value, (tuple, list)):
        return [_jsonable(v) for v in value]
    return str(value)


def exif_pillow(images: list, out: Path) -> dict:
    from PIL import ExifTags, Image
    by_name = {}
    for p in images:
        key = str(p.relative_to(out))
        try:
            with Image.open(p) as img:
                exif = img.getexif()
                tags = {ExifTags.TAGS.get(k, str(k)): _jsonable(v) for k, v in exif.items()}
                gps = exif.get_ifd(0x8825)
                if gps:
                    tags["GPSInfo"] = {ExifTags.GPSTAGS.get(k, str(k)): _jsonable(v)
                                       for k, v in gps.items()}
                by_name[key] = tags
        except Exception as exc:  # noqa: BLE001 -- per-image error recorded, not swallowed
            by_name[key] = {"error": f"{type(exc).__name__}: {exc}"}
    return by_name


def write_exif(out: Path, ext: str) -> dict:
    images = find_images(out, ext)
    if not images:
        return {}
    if shutil.which("exiftool"):
        tool, data = "exiftool", exif_exiftool(images, out)
    else:
        try:
            import PIL  # noqa: F401
        except ImportError:
            tool, data = "unavailable", {}
        else:
            tool, data = "pillow", exif_pillow(images, out)
    note = ("exif: unavailable (install exiftool or Pillow)" if tool == "unavailable"
            else f"exif via {tool}")
    (out / "exif.json").write_text(json.dumps(
        {"generator": SIDE_GENERATOR, "tool": tool, "note": note,
         "images": [str(p.relative_to(out)) for p in images], "exif": data},
        indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    return {"exif": note if tool == "unavailable" else tool, "exif_images": len(images)}


# ---------- optional model NER (--ner) ----------

NER_LABELS = ["person", "organization", "location", "email", "phone number",
              "crypto wallet", "iban", "passport number", "national id number"]
NER_CHUNK = 1500


def _chunks(text: str, size: int = NER_CHUNK):
    """Yield (offset, chunk) cut on whitespace near `size` so entities are not split."""
    pos = 0
    while pos < len(text):
        end = min(len(text), pos + size)
        if end < len(text):
            cut = text.rfind(" ", pos + size // 2, end)
            end = cut if cut > pos else end
        yield pos, text[pos:end]
        pos = end


def ner_gliner(text: str) -> list:
    from gliner import GLiNER
    model = GLiNER.from_pretrained(os.environ.get("GLINER_MODEL", "urchade/gliner_multi-v2.1"))
    found = []
    for offset, chunk in _chunks(text):
        for ent in model.predict_entities(chunk, NER_LABELS, threshold=0.5):
            found.append({"label": ent["label"], "start": offset + ent["start"],
                          "end": offset + ent["end"], "text": ent["text"],
                          "score": round(float(ent.get("score", 0.0)), 4)})
    return found


def ner_langextract(text: str) -> list:
    import langextract as lx
    key = os.environ.get("LANGEXTRACT_API_KEY") or os.environ.get("GEMINI_API_KEY")
    examples = [lx.data.ExampleData(
        text="Wire from John Smith at Acme Holdings Ltd to IBAN GB82WEST12345698765432.",
        extractions=[
            lx.data.Extraction(extraction_class="person", extraction_text="John Smith"),
            lx.data.Extraction(extraction_class="organization",
                               extraction_text="Acme Holdings Ltd"),
            lx.data.Extraction(extraction_class="iban",
                               extraction_text="GB82WEST12345698765432"),
        ])]
    result = lx.extract(
        text_or_documents=text, examples=examples, api_key=key,
        model_id=os.environ.get("LANGEXTRACT_MODEL", "gemini-2.5-flash"),
        prompt_description=("Extract people, organizations, locations, emails, phone numbers, "
                            "crypto wallets, IBANs and passport / national ID numbers verbatim."))
    found = []
    for ext in result.extractions or []:
        interval = getattr(ext, "char_interval", None)
        if not interval or interval.start_pos is None:
            continue
        found.append({"label": ext.extraction_class, "start": interval.start_pos,
                      "end": interval.end_pos, "text": ext.extraction_text})
    return found


SPACY_PATTERNS = [
    {"label": "IBAN", "pattern": [{"TEXT": {"REGEX": r"^[A-Z]{2}\d{2}[A-Z0-9]{11,30}$"}}]},
    {"label": "CRYPTO_WALLET", "pattern": [{"TEXT": {"REGEX": r"^0x[a-fA-F0-9]{40}$"}}]},
    {"label": "CRYPTO_WALLET",
     "pattern": [{"TEXT": {"REGEX": r"^(bc1[ac-hj-np-z02-9]{6,87}|[13][a-km-zA-HJ-NP-Z1-9]{25,34})$"}}]},
    {"label": "CRYPTO_WALLET", "pattern": [{"TEXT": {"REGEX": r"^T[1-9A-HJ-NP-Za-km-z]{33}$"}}]},
    {"label": "PASSPORT_OR_ID", "pattern": [{"TEXT": {"REGEX": r"^[A-Z]{1,2}\d{6,9}$"}}]},
    {"label": "PASSPORT_OR_ID", "pattern": [{"LOWER": {"IN": ["passport", "id", "ssn", "nin"]}},
                                            {"IS_PUNCT": True, "OP": "?"},
                                            {"TEXT": {"REGEX": r"^[A-Z0-9-]{6,14}$"}}]},
]


def ner_spacy(text: str) -> tuple:
    import spacy
    try:
        nlp = spacy.load("en_core_web_sm")
        model_note = "en_core_web_sm + EntityRuler"
        ruler = nlp.add_pipe("entity_ruler", before="ner")
    except OSError:
        nlp = spacy.blank("en")
        model_note = "blank('en') + EntityRuler (en_core_web_sm not installed: no PERSON/ORG)"
        ruler = nlp.add_pipe("entity_ruler")
    ruler.add_patterns(SPACY_PATTERNS)
    found = []
    for offset, chunk in _chunks(text, 100000):
        for ent in nlp(chunk).ents:
            found.append({"label": ent.label_, "start": offset + ent.start_char,
                          "end": offset + ent.end_char, "text": ent.text})
    return found, model_note


NER_INSTALL = {"gliner": "pip install gliner", "langextract": "pip install langextract",
               "spacy": "pip install spacy (optional: python -m spacy download en_core_web_sm)"}


def _run_ner_backend(backend: str, text: str) -> tuple:
    if backend == "gliner":
        return ner_gliner(text), "GLiNER " + os.environ.get("GLINER_MODEL",
                                                             "urchade/gliner_multi-v2.1")
    if backend == "langextract":
        return ner_langextract(text), "langextract (Gemini; text left this machine)"
    return ner_spacy(text)


def write_ner(out: Path, sources: dict, backend: str) -> dict:
    """Optional NER -> ner.json. Never raises: missing lib / key becomes a status note."""
    record = {"generator": SIDE_GENERATOR, "backend": backend, "status": "ok", "note": "",
              "offset_unit": "python str index (code point)", "entities": [],
              "dropped_misaligned": 0, "evidence_tier": "T3 (model output)"}
    if backend == "langextract" and not (os.environ.get("LANGEXTRACT_API_KEY")
                                         or os.environ.get("GEMINI_API_KEY")):
        record.update(status="skipped",
                      note="langextract needs LANGEXTRACT_API_KEY or GEMINI_API_KEY")
    else:
        try:
            for name, text in sources.items():
                found, model_note = _run_ner_backend(backend, text)
                record["note"] = model_note
                for ent in found:
                    if text[ent["start"]:ent["end"]] != ent["text"]:
                        record["dropped_misaligned"] += 1  # span not verbatim: not citable
                        continue
                    record["entities"].append({**ent, "source_file": name})
        except ImportError as exc:
            record.update(status="unavailable",
                          note=f"{backend} not installed ({exc}): {NER_INSTALL[backend]}")
        except Exception as exc:  # noqa: BLE001 -- model/runtime failure is reported, not fatal
            record.update(status="error", note=f"{type(exc).__name__}: {exc}")
    (out / "ner.json").write_text(json.dumps(record, indent=2, ensure_ascii=False),
                                  encoding="utf-8")
    if record["status"] != "ok":
        print(f"  [ner] {out.name}: {record['status']} -- {record['note']}", file=sys.stderr)
    return {"ner": record["status"], "ner_entities": len(record["entities"])}


def write_side_files(out: Path, ext: str, ner_backend: str | None) -> dict:
    """iocs/entities/exif (+ner). Errors are reported in the manifest and on stderr,
    never raised: the verbatim extraction already on disk must not be lost."""
    stats: dict = {}
    steps = [("iocs/entities", lambda: write_iocs_and_entities(out, read_side_sources(out))),
             ("exif", lambda: write_exif(out, ext))]
    if ner_backend:
        steps.append(("ner", lambda: write_ner(out, read_side_sources(out), ner_backend)))
    for label, step in steps:
        try:
            stats.update(step())
        except Exception as exc:  # noqa: BLE001 -- recorded + printed, extraction continues
            stats.setdefault("side_file_errors", []).append(f"{label}: {exc}")
            print(f"  [warn] {out.name}: {label} side file failed: {exc}", file=sys.stderr)
    return stats


HANDLERS = {
    ".pdf": extract_pdf,
    ".json": extract_json,
    ".xlsx": extract_xlsx,
    ".docx": extract_docx,
    ".doc": extract_doc,
    ".csv": extract_csv,
    ".eml": extract_eml,
    ".md": extract_md,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
    ".tif": extract_image,
    ".tiff": extract_image,
}


# ---------- orchestration ----------

def resolve_paths(args):
    case, case_dir = resolve_case_dir(args.case or "", WORKSPACE)
    intake = Path(args.intake) if args.intake else case_dir / "investigation" / "intake"
    if not intake.is_dir():
        print(f"ERROR: intake dir not found: {intake}", file=sys.stderr)
        sys.exit(1)
    extracted = case_dir / "investigation" / "evidence" / "extracted"
    extracted.mkdir(parents=True, exist_ok=True)
    return case, intake, extracted


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract verbatim text + full OCR from intake files.")
    parser.add_argument("--case", help="case folder name under investigations/")
    parser.add_argument("--intake", help="override intake dir")
    parser.add_argument("--ner", choices=["gliner", "langextract", "spacy"],
                        help="optional model NER backend -> ner.json (lazy, never fatal)")
    args = parser.parse_args()

    case, intake, extracted = resolve_paths(args)
    files = sorted(p for p in intake.iterdir() if p.is_file())
    if not files:
        print(f"ERROR: no files in {intake}", file=sys.stderr)
        return 1

    print(f"Case: {case}")
    print(f"Intake: {intake}  ({len(files)} files)\n")
    print(f"{'FILE':<52} {'TEXT_CH':>8} {'IMGS':>5} {'OCR_KB':>7} {'STATUS':>7}")

    failures = []
    for src in files:
        handler = HANDLERS.get(src.suffix.lower())
        if handler is None:
            print(f"{src.name[:52]:<52} {'-':>8} {'-':>5} {'-':>7} {'SKIP':>7}")
            continue
        out = extracted / src.stem
        try:
            stats = handler(src, out)
            stats.update(write_side_files(out, src.suffix.lower(), args.ner))
            ocr_kb = (out / "ocr.md").stat().st_size / 1024 if (out / "ocr.md").exists() else 0
            stats.update({"file": src.name, "ext": src.suffix.lower(),
                          "ocr_bytes": int(ocr_kb * 1024), "status": "OK",
                          "extractor": EXTRACTOR_NAME,
                          "extracted_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")})
            (out / "manifest.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")
            print(f"{src.name[:52]:<52} {stats['text_chars']:>8} {stats['images']:>5} "
                  f"{ocr_kb:>7.1f} {'OK':>7}")
        except Exception as exc:  # noqa: BLE001 -- report and continue, preserve partial evidence
            failures.append((src.name, str(exc)))
            print(f"{src.name[:52]:<52} {'-':>8} {'-':>5} {'-':>7} {'FAIL':>7}")

    print(f"\nExtracted -> {extracted}")
    if failures:
        print("\nFAILURES:", file=sys.stderr)
        for name, err in failures:
            print(f"  {name}: {err}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
